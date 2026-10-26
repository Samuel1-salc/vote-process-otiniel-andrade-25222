"""Exportação auditável da votação validada para XLSX."""

from __future__ import annotations

import os
import tempfile
from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.quality_report import QualityReport
from src.section_votes import CandidateSectionVotes

DATA_SHEET = "Votos por seção"
METADATA_SHEET = "Metadados"
HEADER_ROW = 6
FIRST_DATA_ROW = HEADER_ROW + 1
HEADERS = (
    "Cidade",
    "Código Município TSE",
    "Zona Eleitoral",
    "Seção",
    "Seções Agregadas",
    "Votos",
)


class XlsxExportError(RuntimeError):
    """A exportação XLSX não atende às invariantes obrigatórias."""


def _safe_text(value: str) -> str:
    """Impede que texto de origem seja interpretado como fórmula."""
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _validate_export_inputs(
    records: list[CandidateSectionVotes], report: QualityReport
) -> None:
    if report.status != "aprovado":
        raise XlsxExportError("O relatório de qualidade não está aprovado.")
    metrics = report.metrics
    calculated = {
        "registros": len(records),
        "municipios": len({record.codigo_municipio_tse for record in records}),
        "agregacoes": sum(bool(record.secoes_agregadas) for record in records),
        "votos": sum(record.votos for record in records),
    }
    expected = {
        "registros": metrics.records,
        "municipios": metrics.municipalities,
        "agregacoes": metrics.aggregated_records,
        "votos": metrics.total_votes,
    }
    if calculated != expected:
        raise XlsxExportError(
            f"Registros e relatório de qualidade divergem: {calculated!r} != {expected!r}."
        )


def _style_data_sheet(sheet, last_row: int, total_votes: int) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_bold = Font(color="FFFFFF", bold=True)

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Votação por seção — Otoniel Andrade (25222)"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = dark_fill
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 26

    sheet["A2"] = "Cargo"
    sheet["B2"] = "Deputado Estadual — TO — Eleições 2026 — 1º turno"
    sheet["A3"] = "Total validado"
    sheet["B3"] = total_votes
    sheet["A4"] = "Total visível após filtros"
    sheet["B4"] = f"=SUBTOTAL(109,F{FIRST_DATA_ROW}:F{last_row})"
    sheet["B3"].number_format = "#,##0"
    sheet["B4"].number_format = "#,##0"
    for row in (2, 3, 4):
        sheet.cell(row, 1).font = Font(bold=True)
        sheet.cell(row, 1).fill = light_fill

    sheet.merge_cells("A5:F5")
    sheet["A5"] = (
        "Seções agregadas representam uma única unidade de apuração; "
        "os votos não foram distribuídos entre as seções."
    )
    sheet["A5"].font = Font(italic=True, color="666666")

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(HEADER_ROW, column, header)
        cell.fill = dark_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:F{last_row}"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 23
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 22
    sheet.column_dimensions["F"].width = 12


def _build_metadata_rows(report: QualityReport) -> list[tuple[str, object]]:
    source = report.source
    candidate = report.candidate
    metrics = report.metrics
    return [
        ("Status da validação", report.status),
        ("Versão do contrato", report.contract_version),
        ("Candidato", candidate["nome"]),
        ("Número", candidate["numero"]),
        ("Cargo", candidate["cargo"]),
        ("UF", candidate["uf"]),
        ("Ano da eleição", candidate["ano_eleicao"]),
        ("Turno", candidate["turno"]),
        ("Registros", metrics.records),
        ("Municípios", metrics.municipalities),
        ("Registros com seção agregada", metrics.aggregated_records),
        ("Total de votos", metrics.total_votes),
        ("Fonte", source.get("fonte", "TSE — Boletim de Urna")),
        ("URL oficial", source.get("url_origem", "")),
        ("Algoritmo do hash", source.get("algoritmo_hash", "")),
        ("SHA-512 calculado", source.get("hash_calculado", "")),
        ("SHA-512 esperado", source.get("hash_esperado", "")),
        ("Data da aquisição", source.get("adquirido_em", "")),
        ("Data do relatório", report.generated_at),
        (
            "Observação",
            "Seções agregadas são preservadas sem replicar ou distribuir votos.",
        ),
    ]


def build_workbook(
    records: Iterable[CandidateSectionVotes], report: QualityReport
) -> Workbook:
    materialized = list(records)
    _validate_export_inputs(materialized, report)

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = DATA_SHEET
    for row_number, record in enumerate(materialized, start=FIRST_DATA_ROW):
        values: tuple[object, ...] = (
            _safe_text(record.cidade),
            _safe_text(record.codigo_municipio_tse),
            _safe_text(record.zona_eleitoral),
            _safe_text(record.secao),
            _safe_text(", ".join(record.secoes_agregadas)),
            record.votos,
        )
        for column, value in enumerate(values, start=1):
            cell = data_sheet.cell(row_number, column, value)
            if column < 6:
                cell.number_format = "@"
            else:
                cell.number_format = "#,##0"

    last_row = FIRST_DATA_ROW + len(materialized) - 1
    _style_data_sheet(data_sheet, last_row, report.metrics.total_votes)
    table = Table(displayName="VotosPorSecao", ref=f"A{HEADER_ROW}:F{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    data_sheet.add_table(table)

    metadata_sheet = workbook.create_sheet(METADATA_SHEET)
    metadata_sheet.append(("Campo", "Valor"))
    for label, value in _build_metadata_rows(report):
        metadata_sheet.append((_safe_text(label), _safe_text(value) if isinstance(value, str) else value))
    for cell in metadata_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    metadata_sheet.freeze_panes = "A2"
    metadata_sheet.column_dimensions["A"].width = 34
    metadata_sheet.column_dimensions["B"].width = 110
    metadata_sheet.auto_filter.ref = f"A1:B{metadata_sheet.max_row}"
    for row in metadata_sheet.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def _verify_saved_workbook(path: Path, expected_records: int, expected_votes: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != [DATA_SHEET, METADATA_SHEET]:
            raise XlsxExportError(f"Abas inesperadas no XLSX: {workbook.sheetnames!r}")
        sheet = workbook[DATA_SHEET]
        last_row = FIRST_DATA_ROW + expected_records - 1
        if sheet.max_row != last_row:
            raise XlsxExportError(
                f"Quantidade de linhas inesperada: {sheet.max_row}, esperado {last_row}."
            )
        if tuple(sheet.cell(HEADER_ROW, index).value for index in range(1, 7)) != HEADERS:
            raise XlsxExportError("Cabeçalho XLSX incompatível.")
        if sheet["B3"].value != expected_votes:
            raise XlsxExportError("Total validado não foi preservado no XLSX.")
        expected_formula = f"=SUBTOTAL(109,F{FIRST_DATA_ROW}:F{last_row})"
        if sheet["B4"].value != expected_formula:
            raise XlsxExportError("Fórmula de total dinâmico incompatível.")
    finally:
        workbook.close()


def export_workbook_atomic(
    records: Iterable[CandidateSectionVotes], report: QualityReport, destination: Path
) -> None:
    materialized = list(records)
    workbook = build_workbook(materialized, report)
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=".xlsx", dir=destination.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        workbook.close()
        _verify_saved_workbook(
            temporary,
            expected_records=report.metrics.records,
            expected_votes=report.metrics.total_votes,
        )
        os.replace(temporary, destination)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
