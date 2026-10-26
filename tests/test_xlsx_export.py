from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from src.quality_report import validate_candidate_votes
from src.section_votes import CandidateSectionVotes
from src.xlsx_export import (
    DATA_SHEET,
    FIRST_DATA_ROW,
    HEADER_ROW,
    METADATA_SHEET,
    XlsxExportError,
    export_workbook_atomic,
)


def record(**changes: object) -> CandidateSectionVotes:
    base = CandidateSectionVotes(
        cidade="NATIVIDADE",
        codigo_municipio_tse="94811",
        zona_eleitoral="19",
        secao="81",
        secoes_agregadas=(),
        votos=4,
        numero_candidato="25222",
        nome_candidato="OTONIEL ANDRADE",
        cargo="Deputado Estadual",
        uf="TO",
        ano_eleicao=2026,
        turno=1,
    )
    return replace(base, **changes)


class XlsxExportTest(unittest.TestCase):
    def _export(self, records: list[CandidateSectionVotes], destination: Path) -> None:
        report = validate_candidate_votes(
            records,
            source={
                "versao_contrato": "1.1.0",
                "url_origem": "https://example.invalid/bu.zip",
                "algoritmo_hash": "SHA-512",
                "hash_calculado": "abc123",
            },
            generated_at="2026-10-26T00:00:00Z",
        )
        export_workbook_atomic(records, report, destination)

    def test_workbook_structure_filters_formula_and_types(self) -> None:
        records = [record(), record(secao="82", votos=3, secoes_agregadas=("83",))]
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "result.xlsx"
            self._export(records, destination)
            workbook = load_workbook(destination, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, [DATA_SHEET, METADATA_SHEET])
                sheet = workbook[DATA_SHEET]
                self.assertEqual(sheet.freeze_panes, f"A{FIRST_DATA_ROW}")
                self.assertEqual(sheet.auto_filter.ref, f"A{HEADER_ROW}:F8")
                self.assertEqual(sheet["B4"].value, "=SUBTOTAL(109,F7:F8)")
                self.assertEqual(sheet["C7"].data_type, "s")
                self.assertEqual(sheet["D7"].data_type, "s")
                self.assertEqual(sheet["F7"].data_type, "n")
                self.assertEqual(sheet["E8"].value, "83")
                self.assertEqual(sheet["B3"].value, 7)
            finally:
                workbook.close()

    def test_metadata_contains_source_contract_and_total(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "result.xlsx"
            self._export([record()], destination)
            workbook = load_workbook(destination, data_only=False)
            try:
                rows = dict(workbook[METADATA_SHEET].iter_rows(min_row=2, values_only=True))
                self.assertEqual(rows["Versão do contrato"], "1.1.0")
                self.assertEqual(rows["Algoritmo do hash"], "SHA-512")
                self.assertEqual(rows["Total de votos"], 4)
                self.assertEqual(rows["URL oficial"], "https://example.invalid/bu.zip")
            finally:
                workbook.close()

    def test_formula_like_text_is_neutralized(self) -> None:
        unsafe = record(cidade="=SUM(A1:A2)")
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "result.xlsx"
            self._export([unsafe], destination)
            workbook = load_workbook(destination, data_only=False)
            try:
                self.assertEqual(workbook[DATA_SHEET]["A7"].value, "'=SUM(A1:A2)")
                self.assertEqual(workbook[DATA_SHEET]["A7"].data_type, "s")
            finally:
                workbook.close()

    def test_rejects_unapproved_report_without_partial_file(self) -> None:
        records = [record()]
        report = validate_candidate_votes(records)
        report = replace(report, status="reprovado")
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "result.xlsx"
            with self.assertRaises(XlsxExportError):
                export_workbook_atomic(records, report, destination)
            self.assertFalse(destination.exists())

    def test_rejects_metrics_mismatch_without_partial_file(self) -> None:
        first = record()
        report = validate_candidate_votes([first])
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "result.xlsx"
            with self.assertRaises(XlsxExportError):
                export_workbook_atomic([first, record(secao="82")], report, destination)
            self.assertFalse(destination.exists())


if __name__ == "__main__":
    unittest.main()
